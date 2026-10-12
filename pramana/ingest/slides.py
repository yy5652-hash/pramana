"""PowerPoint decks -> one citable unit per slide, plus a rendered preview of each slide.

Two things a plain text extractor misses, both handled here:

* PowerPoint stores an equation as Office Math inside an `mc:AlternateContent` block, with a
  pre-rendered picture of the same box as fallback. python-pptx skips those blocks, so a
  naive extractor silently drops every formula on a slide. We read the block ourselves: the
  fallback picture goes into the preview (and to the vision model, which returns LaTeX), and
  the linearised Office Math text is kept as a no-vision fallback.
* The preview is drawn with Pillow from the shapes themselves, so no office suite is needed.
"""
from __future__ import annotations

import io
import unicodedata
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Callable

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from . import vision
from .common import file_digest

NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
}
PREVIEW_WIDTH = 1280
FONT_PATHS = (
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/Library/Fonts/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "C:/Windows/Fonts/arial.ttf",
)
INK = (28, 30, 36)


def _font(size: int) -> ImageFont.FreeTypeFont:
    for candidate in FONT_PATHS:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default(size=size)


def _plain(text: str) -> str:
    """Office Math uses mathematical-alphanumeric code points; fold them to ordinary letters."""
    return unicodedata.normalize("NFKC", text).replace("\x0b", "\n")


def _paragraph_text(paragraph) -> str:
    """Text of an `a:p`, with any Office Math it contains linearised in place."""
    parts = []
    for node in paragraph.iter():
        if node.tag == f"{{{NS['a']}}}t" or node.tag == f"{{{NS['m']}}}t":
            parts.append(node.text or "")
        elif node.tag == f"{{{NS['a']}}}br":
            parts.append("\n")
    return _plain("".join(parts)).strip()


def _frame(element) -> tuple[int, int, int, int] | None:
    offset = element.find(".//a:xfrm/a:off", NS)
    extent = element.find(".//a:xfrm/a:ext", NS)
    if offset is None or extent is None:
        return None
    return int(offset.get("x")), int(offset.get("y")), int(extent.get("cx")), int(extent.get("cy"))


def _slide_items(slide) -> list[dict]:
    """Every drawable thing on a slide, in z-order: text boxes, pictures, tables, maths boxes."""
    items: list[dict] = []
    shapes_by_element = {id(shape._element): shape for shape in slide.shapes}
    for element in slide.shapes._spTree:
        if element.tag == f"{{{NS['mc']}}}AlternateContent":
            choice = element.find("mc:Choice", NS)
            fallback = element.find("mc:Fallback", NS)
            source = choice if choice is not None else fallback
            text = "\n".join(t for t in (_paragraph_text(p) for p in source.iterfind(".//a:p", NS)) if t)
            has_math = source.find(f".//{{{NS['m']}}}oMath") is not None
            blob = None
            blip = fallback.find(".//a:blip", NS) if fallback is not None else None
            if blip is not None and blip.get(f"{{{NS['r']}}}embed"):
                blob = slide.part.related_part(blip.get(f"{{{NS['r']}}}embed")).blob
            items.append({"kind": "math" if has_math else "text", "frame": _frame(source), "text": text,
                          "image": blob, "role": "body", "size": None})
            continue
        shape = shapes_by_element.get(id(element))
        if shape is None:
            continue
        frame = (shape.left, shape.top, shape.width, shape.height) if shape.width else None
        # Pictures come as p:pic, as picture placeholders, or as shapes with a picture fill;
        # what they share is an embedded `a:blip`.
        blip = element.find(".//a:blip", NS)
        if blip is not None and blip.get(f"{{{NS['r']}}}embed"):
            blob = slide.part.related_part(blip.get(f"{{{NS['r']}}}embed")).blob
            items.append({"kind": "picture", "frame": frame, "text": "", "image": blob,
                          "role": "picture", "size": None})
        elif getattr(shape, "has_table", False) and shape.has_table:
            rows = [" | ".join(_plain(cell.text).strip() for cell in row.cells) for row in shape.table.rows]
            items.append({"kind": "table", "frame": frame, "text": "\n".join(rows), "image": None,
                          "role": "body", "size": None})
        elif shape.has_text_frame and shape.text_frame.text.strip():
            role = "body"
            if shape.is_placeholder and shape.placeholder_format.type in (
                PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE
            ):
                role = "title"
            size = next(
                (run.font.size.pt for para in shape.text_frame.paragraphs for run in para.runs if run.font.size),
                None,
            )
            text = "\n".join(t for t in (_paragraph_text(p._p) for p in shape.text_frame.paragraphs) if t)
            items.append({"kind": "text", "frame": frame, "text": text, "image": None, "role": role, "size": size})
    return items


def _wrap(draw: ImageDraw.ImageDraw, text: str, font, width: int) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        current = ""
        for word in paragraph.split():
            trial = f"{current} {word}".strip()
            if draw.textlength(trial, font=font) <= width or not current:
                current = trial
            else:
                lines.append(current)
                current = word
        lines.append(current)
    return lines


def render_slide(items: list[dict], slide_width: int, slide_height: int) -> Image.Image:
    scale = PREVIEW_WIDTH / slide_width
    canvas = Image.new("RGB", (PREVIEW_WIDTH, round(slide_height * scale)), "white")
    draw = ImageDraw.Draw(canvas)
    for item in items:
        if not item["frame"]:
            continue
        x, y, w, h = (round(v * scale) for v in item["frame"])
        if w <= 2 or h <= 2:
            continue
        if item["image"]:
            try:
                picture = Image.open(io.BytesIO(item["image"])).convert("RGBA").resize((w, h))
            except Exception:  # an exotic format (EMF/WMF) Pillow cannot decode
                continue
            canvas.paste(picture, (x, y), picture)
            continue
        if not item["text"]:
            continue
        points = item["size"] or (30 if item["role"] == "title" else 20)
        size = max(12, round(points * scale * 12700))  # points -> EMU -> pixels
        while True:  # shrink until the text fits its box, as PowerPoint's autofit does
            font = _font(size)
            lines = _wrap(draw, item["text"], font, w - 12)
            if len(lines) * size * 1.25 <= h + size or size <= 13:
                break
            size -= 2
        cursor = y + 4
        for line in lines:
            draw.text((x + 6, cursor), line, font=font, fill=INK)
            cursor += round(size * 1.25)
    return canvas


def ingest_slides(
    source: dict,
    path: Path,
    asset_dir: Path,
    progress: Callable[[str], None] = print,
    use_vision: bool = True,
    workers: int = 4,
) -> list[dict]:
    deck = Presentation(str(path))
    digest = file_digest(path)
    asset_dir.mkdir(parents=True, exist_ok=True)
    slide_area = deck.slide_width * deck.slide_height
    slides = []
    for number, slide in enumerate(deck.slides, 1):
        items = _slide_items(slide)
        image_path = asset_dir / f"s{number:03d}.jpg"
        if not image_path.exists():
            render_slide(items, deck.slide_width, deck.slide_height).save(image_path, quality=85)
        title = next((i["text"] for i in items if i["role"] == "title"), "")
        body = [i["text"] for i in items if i["role"] != "title" and i["text"]]
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        # A slide is worth a vision call when it carries a real picture or a rendered equation.
        visual = any(
            (i["kind"] == "math" and i["image"])
            or (i["kind"] == "picture" and i["frame"] and i["frame"][2] * i["frame"][3] > 0.04 * slide_area)
            for i in items
        )
        slides.append({"number": number, "title": " ".join(title.split()), "body": body, "notes": notes,
                       "image": image_path, "visual": visual})

    seen: dict[int, dict] = {}
    if use_vision:
        def read(slide: dict) -> tuple[int, dict]:
            text = "\n".join([slide["title"], *slide["body"]])
            return slide["number"], vision.read_slide(
                slide["image"], text, cache_key=["slide", vision.SLIDE_PROMPT_VERSION, digest, slide["number"]]
            )

        with ThreadPoolExecutor(max_workers=workers) as pool:
            seen = dict(pool.map(read, [s for s in slides if s["visual"]]))
        progress(f"{source['id']}: vision read {len(seen)}/{len(slides)} slides")

    units = []
    for slide in slides:
        parts = [f"### {slide['title']}"] if slide["title"] else []
        parts.extend(slide["body"])
        reading = seen.get(slide["number"])
        if reading:
            if reading.get("equations"):
                parts.append("Equations on the slide: " + "; ".join(f"${e}$" for e in reading["equations"]))
            if reading.get("has_visual") and reading.get("description"):
                parts.append("Figure on the slide: " + reading["description"])
        if slide["notes"]:
            parts.append("Speaker notes: " + slide["notes"])
        text = "\n\n".join(parts)
        if len(text.split()) < 6:
            continue
        units.append({
            "id": f"{source['id']}:s{slide['number']}",
            "source_id": source["id"],
            "kind": "slide",
            "modality": "vision" if reading and (reading.get("equations") or reading.get("has_visual")) else "text",
            "locator": {"slide": slide["number"]},
            "section": slide["title"] or source["title"],
            "text": text,
            "asset": f"{source['id']}/{slide['image'].name}",
        })
    for order, unit in enumerate(units):
        unit["ord"] = order
    return units
